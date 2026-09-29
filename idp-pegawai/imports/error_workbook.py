"""Safe, self-contained error-workbook artifact for the pegawai retry loop."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from copy import copy
from io import BytesIO
from typing import Any
from uuid import uuid4

from openpyxl import Workbook

PROFILE = "pegawai_error_retry_v1"
HEADERS = ("nip", "nama_lengkap", "email_dinas", "email_google", "org_unit_code", "status_pegawai")
SHEETS = ("instructions", "summary", "pegawai", "diagnostics", "reference_status", "reference_org_units", "metadata")
ERROR_HEADERS = ("error_field", "error_message")
STATUS_VALUES = ("aktif", "mutasi_keluar", "pensiun", "nonaktif", "meninggal")


def _validation_messages(issues):
    grouped = {}
    for issue in issues:
        grouped.setdefault(issue.row_number, []).append(f"{issue.field or '-'}: {issue.message or issue.code}")
    return {row: "; ".join(messages) for row, messages in grouped.items()} 


def build_import_template(*, org_units=()):
    """Build an editable XLSX template with safe references and Excel dropdowns."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    instructions = workbook.create_sheet("instructions")
    instructions.append(["Template Import Pegawai"])
    instructions.append(["Isi sheet pegawai mulai dari baris kedua, lalu simpan sebagai XLSX."])
    instructions.append(["Kolom error_field dan error_message hanya muncul pada file error; boleh dibiarkan kosong saat upload template."])
    reference = workbook.create_sheet("reference")
    reference.append(["status_pegawai", "org_unit_code"])
    codes = [str(code) for code in org_units if code]
    for index in range(max(len(STATUS_VALUES), len(codes))):
        reference.append([STATUS_VALUES[index] if index < len(STATUS_VALUES) else None, codes[index] if index < len(codes) else None])
    sheet = workbook.create_sheet("pegawai")
    sheet.append(list(HEADERS))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{chr(64 + len(HEADERS))}1"
    from openpyxl.worksheet.datavalidation import DataValidation
    status_validation = DataValidation(type="list", formula1="=reference!$A$2:$A$6", allow_blank=False)
    if codes:
        unit_validation = DataValidation(type="list", formula1=f"=reference!$B$2:$B${len(codes) + 1}", allow_blank=False)
        sheet.add_data_validation(unit_validation)
        unit_validation.add(f"E2:E1000")
    sheet.add_data_validation(status_validation)
    status_validation.add("F2:F1000")
    for cell in sheet[1]:
        font = copy(cell.font); font.bold = True; cell.font = font
    reference.sheet_state = "hidden"
    stream = BytesIO(); workbook.save(stream)
    return stream.getvalue() 



def escape_formula(value: Any) -> Any:
    """Escape spreadsheet-leading formula characters for display only."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


@dataclass(frozen=True)
class ErrorWorkbookArtifact:
    profile: str
    sheets: dict[str, list[dict[str, Any]]]
    display_sheets: dict[str, list[dict[str, Any]]]
    metadata: dict[str, Any]
    content: bytes


def build_error_workbook(rows, validation, *, parent_batch_uuid, root_batch_uuid, org_units=() ) -> ErrorWorkbookArtifact:
    """Build all-row retry data while retaining canonical values separately from display values."""
    canonical = [{h: row.get(h) for h in HEADERS} for row in rows if any(row.get(h) not in (None, "") for h in HEADERS)]
    issues = getattr(validation, "issues", [])
    issue_messages = _validation_messages(issues)
    issue_fields = {}
    for issue in issues:
        issue_fields.setdefault(issue.row_number, []).append(issue.field or "-")
    display = []
    for index, row in enumerate(canonical, 1):
        item = {h: escape_formula(row[h]) for h in HEADERS}
        item["error_field"] = ", ".join(issue_fields.get(index, []))
        item["error_message"] = issue_messages.get(index, "")
        display.append(item)
    diagnostics = [{"code": i.code, "severity": i.severity, "row_number": i.row_number, "field": i.field, "message": i.message} for i in issues]
    metadata = {"profile": PROFILE, "version": "1", "parent_batch_uuid": str(parent_batch_uuid), "root_batch_uuid": str(root_batch_uuid), "retry_lineage_token": uuid4().hex, "created_at": datetime.now(timezone.utc).isoformat()}
    sheets = {"instructions": [{"text": "Perbaiki data pada sheet pegawai lalu upload kembali."}], "summary": [{"retry_data_rows": len(canonical), "error_diagnostics": sum(i.severity in {"error", "fatal"} for i in issues), "warning_diagnostics": sum(i.severity == "warning" for i in issues)}], "pegawai": canonical, "diagnostics": diagnostics, "reference_status": [{"status": s} for s in ("aktif", "mutasi_keluar", "pensiun", "nonaktif", "meninggal")], "reference_org_units": [], "metadata": [metadata]}
    display_sheets = dict(sheets)
    display_sheets["pegawai"] = display
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in SHEETS:
        worksheet = workbook.create_sheet(title=name)
        data = display_sheets[name]
        headers = list(data[0]) if data else ["value"]
        worksheet.append(headers)
        for row in data:
            worksheet.append([row.get(header) for header in headers])
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            font = copy(cell.font)
            font.bold = True
            cell.font = font
    reference = workbook.create_sheet("reference")
    reference.append(["status_pegawai", "org_unit_code"])
    codes = [str(code) for code in org_units if code]
    for index in range(max(len(STATUS_VALUES), len(codes))):
        reference.append([STATUS_VALUES[index] if index < len(STATUS_VALUES) else None, codes[index] if index < len(codes) else None])
    reference.sheet_state = "hidden"
    from openpyxl.worksheet.datavalidation import DataValidation
    pegawai = workbook["pegawai"]
    status_validation = DataValidation(type="list", formula1="=reference!$A$2:$A$6", allow_blank=False)
    pegawai.add_data_validation(status_validation)
    status_validation.add("F2:F1000")
    if codes:
        unit_validation = DataValidation(type="list", formula1=f"=reference!$B$2:$B${len(codes) + 1}", allow_blank=False)
        pegawai.add_data_validation(unit_validation)
        unit_validation.add("E2:E1000")
    stream = BytesIO()
    workbook.save(stream)
    return ErrorWorkbookArtifact(PROFILE, sheets, display_sheets, metadata, stream.getvalue())


__all__ = ["ErrorWorkbookArtifact", "build_error_workbook", "escape_formula"]
